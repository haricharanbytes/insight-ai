"""Load CSV / Excel bytes into a DataFrame, with validation and metadata."""
import csv
import io
import logging
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import load_workbook

from .validation import (
    EmptyFileError,
    MalformedFileError,
    SheetSelectionRequired,
    TooManyRowsError,
    normalize_columns,
    validate_upload,
)

logger = logging.getLogger("insightai.loader")

_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")


@dataclass
class LoadResult:
    df: pd.DataFrame
    filename: str
    file_type: str  # "csv" or "xlsx"
    size_bytes: int
    encoding: str | None = None
    delimiter: str | None = None
    sheet_name: str | None = None
    sheet_names: list[str] = field(default_factory=list)
    renamed_columns: dict[str, str] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def get_excel_sheet_names(content: bytes) -> list[str]:
    """Return sheet names of an .xlsx workbook (without reading the data)."""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception as exc:  # openpyxl raises many different error types
        raise MalformedFileError(f"Could not open the Excel workbook: {exc}") from exc


def _decode_text(content: bytes) -> tuple[str, str]:
    for enc in _ENCODINGS:
        try:
            return content.decode(enc), enc
        except UnicodeDecodeError:
            continue
    # latin-1 never fails, so this is unreachable; kept for safety.
    raise MalformedFileError("Could not decode the file as text.")


def _detect_delimiter(text: str) -> str:
    sample = text[:20000]
    if "\n" in sample:
        sample = sample[: sample.rfind("\n")]  # avoid a half-cut last line
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _read_csv(content: bytes, max_rows: int) -> tuple[pd.DataFrame, str, str]:
    text, encoding = _decode_text(content)
    delimiter = _detect_delimiter(text)
    try:
        df = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            nrows=max_rows + 1,  # +1 lets us detect "too many rows"
            low_memory=False,
        )
    except pd.errors.EmptyDataError as exc:
        raise EmptyFileError("The CSV file contains no data.") from exc
    except (pd.errors.ParserError, ValueError) as exc:
        raise MalformedFileError(f"Could not parse the CSV file: {exc}") from exc
    return df, encoding, delimiter


def _read_excel(content: bytes, sheet_name: str, max_rows: int) -> pd.DataFrame:
    try:
        return pd.read_excel(
            io.BytesIO(content),
            sheet_name=sheet_name,
            nrows=max_rows + 1,
            engine="openpyxl",
        )
    except ValueError as exc:
        raise MalformedFileError(f"Could not read sheet '{sheet_name}': {exc}") from exc
    except Exception as exc:
        raise MalformedFileError(f"Could not read the Excel file: {exc}") from exc


def load_dataset(
    filename: str,
    content: bytes,
    max_bytes: int,
    max_rows: int,
    sheet_name: str | None = None,
) -> LoadResult:
    """Validate and load an uploaded file.

    Raises a DataLoadError subclass for any user-facing problem.
    For multi-sheet workbooks with no sheet_name given, raises
    SheetSelectionRequired carrying the sheet list.
    """
    safe_name, ext = validate_upload(filename, content, max_bytes)
    warnings: list[str] = []
    encoding = delimiter = None
    sheet_names: list[str] = []
    chosen_sheet: str | None = None

    if ext == ".csv":
        df, encoding, delimiter = _read_csv(content, max_rows)
        if encoding != "utf-8-sig":
            warnings.append(
                f"File was not UTF-8; decoded as {encoding}. "
                "Check that special characters look correct."
            )
        file_type = "csv"
    else:
        sheet_names = get_excel_sheet_names(content)
        if not sheet_names:
            raise EmptyFileError("The workbook has no sheets.")
        if sheet_name is None:
            if len(sheet_names) > 1:
                raise SheetSelectionRequired(sheet_names)
            sheet_name = sheet_names[0]
        elif sheet_name not in sheet_names:
            raise MalformedFileError(
                f"Sheet '{sheet_name}' not found. Available sheets: {', '.join(sheet_names)}."
            )
        chosen_sheet = sheet_name
        df = _read_excel(content, sheet_name, max_rows)
        file_type = "xlsx"

    # Shape checks
    if df.shape[1] == 0 or df.shape[0] == 0:
        raise EmptyFileError("The file has no data rows or no columns.")
    if df.shape[0] > max_rows:
        raise TooManyRowsError(
            f"The dataset has more than {max_rows:,} rows, which is the configured limit."
        )

    # Safe column names (values are never modified)
    new_cols, renames = normalize_columns(df.columns)
    df.columns = new_cols
    if renames:
        warnings.append(
            f"{len(renames)} column name(s) were adjusted (blank, duplicate, or extra spaces)."
        )

    logger.info(
        "Loaded %s: %d rows x %d cols (%s)",
        safe_name, df.shape[0], df.shape[1], file_type,
    )  # row/column counts only, never cell values

    return LoadResult(
        df=df,
        filename=safe_name,
        file_type=file_type,
        size_bytes=len(content),
        encoding=encoding,
        delimiter=delimiter,
        sheet_name=chosen_sheet,
        sheet_names=sheet_names,
        renamed_columns=renames,
        warnings=warnings,
    )