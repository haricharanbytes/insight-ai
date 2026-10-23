"""Render the EDA report model as an .xlsx workbook. Text cells are protected against formula injection."""
import io
import json
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from src.analysis.common import to_py
from src.reports.exporter import safe_cell

HEADER_FILL = PatternFill("solid", fgColor="E4EFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1E3A5F")
BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_WIDTH = 50
MAX_ROWS_PER_SHEET = 5000


def _cell(v):
    if isinstance(v, str):
        return safe_cell(v)
    if isinstance(v, bool) or v is None or isinstance(v, (int, float)):
        return v
    return safe_cell(str(v))


def _sheet_name(wb: Workbook, base: str) -> str:
    name = BAD_SHEET_CHARS.sub(" ", base).strip().strip("'")[:31] or "Sheet"
    candidate, n = name, 2
    existing = {s.lower() for s in wb.sheetnames}
    while candidate.lower() in existing:
        suffix = f" {n}"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    return candidate


def _widths(ws) -> None:
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(MAX_WIDTH, max(10, longest + 2))


def _write_rows(ws, start: int, columns: list, rows: list) -> int:
    for j, c in enumerate(columns, 1):
        cell = ws.cell(row=start, column=j, value=_cell(str(c)))
        cell.font, cell.fill = BOLD, HEADER_FILL
    for i, row in enumerate(rows[:MAX_ROWS_PER_SHEET], start + 1):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=_cell(v))
    return start + 1 + min(len(rows), MAX_ROWS_PER_SHEET)


def _block(ws, start: int, heading: str, lines: list[str]) -> int:
    ws.cell(row=start, column=1, value=_cell(heading)).font = BOLD
    r = start + 1
    for line in lines:
        c = ws.cell(row=r, column=1, value=_cell(line))
        c.alignment = Alignment(wrap_text=False)
        r += 1
    return r + 1


def _add_chart(ws, section_key: str, table: dict, header_row: int, last_row: int, anchor_col: int):
    cols = table["columns"]
    if last_row <= header_row:
        return
    if section_key == "categories" and "row_count" in cols:
        chart = BarChart()
        chart.type = "bar"
        chart.title = f"{cols[0]}: rows per value"
        chart.add_data(Reference(ws, min_col=cols.index("row_count") + 1, min_row=header_row, max_row=last_row),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
    elif section_key == "trends" and "value" in cols and cols[0] == "period":
        chart = LineChart()
        chart.title = "Value per period"
        chart.add_data(Reference(ws, min_col=cols.index("value") + 1, min_row=header_row, max_row=last_row),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
    else:
        return
    chart.height, chart.width = 8, 16
    chart.legend = None
    ws.add_chart(chart, f"{ws.cell(row=header_row, column=anchor_col).column_letter}{header_row}")


def render_eda_excel(report: dict) -> bytes:
    meta = report["meta"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value=_cell(meta["title"])).font = TITLE
    info = [f"Dataset: {meta['dataset']}", f"Rows: {meta['rows']:,}   Columns: {meta['columns']}",
            f"Generated: {meta['generated_at']}"]
    if meta["author"]:
        info.insert(0, f"Author: {meta['author']}")
    if meta["filters"]:
        info.append("Filters: " + json.dumps(meta["filters"], default=to_py))
    r = _block(ws, 3, "Report details", info)
    src = report.get("summary_source", "rules")
    r = _block(ws, r, "Executive summary (" + ("observed, built from calculated results" if src == "rules"
                                               else "AI-written, numbers checked against calculations") + ")",
               report["summary"])
    skipped = [s["title"] for s in report["sections"] if s["ok"] and s["charts"]]
    if skipped:
        r = _block(ws, r, "Charts", ["Native charts are included for category breakdowns and trends. Other charts "
                                     "(histograms, correlation heatmap, scatter) are in the HTML and PDF reports; "
                                     "their tables are in this workbook."])
    failed = [f"{s['title']}: {'; '.join(s['warnings'])}" for s in report["sections"] if not s["ok"]]
    if failed:
        r = _block(ws, r, "Sections not available", failed)
    if report["warnings"]:
        r = _block(ws, r, "Warnings", report["warnings"])
    _block(ws, r, "Notes", report["notes"])
    _widths(ws)
    ws.column_dimensions["A"].width = 100

    for sec in report["sections"]:
        if not sec["ok"]:
            continue
        for k, t in enumerate(sec["tables"]):
            if sec["key"] == "overview":
                name = "Columns"
            elif sec["key"] == "quality":
                name = "Data quality"
            else:
                name = f"{sec['title'].split()[0]} - {t['title']}"
            sh = wb.create_sheet(_sheet_name(wb, name))
            row = 1
            if k == 0:
                lines = list(sec["findings"]) + ([f"Caveat: {sec['caveat']}"] if sec["caveat"] else []) \
                    + [f"Warning: {w}" for w in sec["warnings"]]
                row = _block(sh, 1, sec["title"], lines) if lines else 1
            header_row = row
            last = _write_rows(sh, header_row, t["columns"], t["rows"]) - 1
            _add_chart(sh, sec["key"], t, header_row, last, len(t["columns"]) + 2)
            _widths(sh)
            sh.column_dimensions["A"].width = max(sh.column_dimensions["A"].width or 10, 24)
            sh.freeze_panes = sh.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()