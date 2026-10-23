import io

import pandas as pd
from openpyxl import load_workbook

from src.reports.eda import build_eda_report
from src.reports.eda_excel import _sheet_name, render_eda_excel

N = 120
DF = pd.DataFrame({
    "Region": ["East", "West", "North", "South"] * 30,
    "Segment": ["A", "B", "C"] * 40,
    "Sales": [100.0 + i * 3 + (i % 7) for i in range(N)],
    "Profit": [10.0 - (i % 11) * 2 for i in range(N)],
    "Order Date": [f"{2019 + i // 40}-{(i % 12) + 1:02d}-10" for i in range(N)],
})


def wb_of(df=DF, **opts):
    return load_workbook(io.BytesIO(render_eda_excel(build_eda_report(df, "d.csv", opts or None))))


def test_workbook_has_expected_sheets_and_summary():
    wb = wb_of(title="Q report", author="Asha")
    assert wb.sheetnames[0] == "Summary"
    assert "Columns" in wb.sheetnames
    assert any(n.startswith("Categories - ") for n in wb.sheetnames)
    ws = wb["Summary"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Q report" in text and "Author: Asha" in text and "This report contains no recommendations." in text


def test_numbers_stay_numeric_and_match_the_data():
    wb = wb_of()
    ws = next(wb[n] for n in wb.sheetnames if n.startswith("Numeric - "))
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    header = next(i for i, r in enumerate(rows) if r and r[0] == "Column")
    sales = next(r for r in rows[header + 1:] if r[0] == "Sales")
    assert isinstance(sales[1], float) and abs(sales[1] - DF["Sales"].mean()) < 1e-9
    profit = next(r for r in rows[header + 1:] if r[0] == "Profit")
    assert profit[4] < 0                                  # negative minimum stays a negative number


def test_formula_injection_is_neutralised():
    df = DF.assign(Region=["=HYPERLINK(\"http://x\")", "+cmd", "@sum", "-1+1"] * 30)
    wb = wb_of(df)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not c.value.startswith(("=", "+", "@")), (ws.title, c.value)
                    assert not (c.value.startswith("-") and len(c.value) > 1 and not c.value[1:2].isdigit()), c.value


def test_native_charts_exist_for_categories_and_trends():
    wb = wb_of()
    charts = {ws.title: len(ws._charts) for ws in wb.worksheets}
    assert any(n.startswith("Categories - ") and k == 1 for n, k in charts.items())
    assert any(n.startswith("Trends - ") and k == 1 for n, k in charts.items())


def test_sheet_names_are_valid_and_unique():
    wb = load_workbook(io.BytesIO(render_eda_excel(build_eda_report(DF, "d.csv"))))
    assert len(set(n.lower() for n in wb.sheetnames)) == len(wb.sheetnames)
    assert all(len(n) <= 31 for n in wb.sheetnames)
    from openpyxl import Workbook
    w = Workbook()
    assert _sheet_name(w, "a/b*c?[x]:y" + "z" * 40) == "a b c  x  y" + "z" * 20
    assert _sheet_name(w, "Sheet") != "Sheet"             # collides with the default sheet


def test_failed_sections_are_listed_not_given_sheets():
    cats = pd.DataFrame({"g": list("aabbccaabb") * 5, "h": list("xyxyxyxyxy") * 5})
    wb = wb_of(cats)
    assert not any(n.startswith(("Numeric", "Trends")) for n in wb.sheetnames)
    text = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert "Sections not available" in text